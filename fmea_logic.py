"""Core logic for the FMEA + Pareto analyzer.

Only pandas and openpyxl are used here (no Streamlit, no Plotly), so everything in
this module can be tested on its own.
"""
from __future__ import annotations

import io

import pandas as pd

TEXT_COLS = ["Failure Mode", "Effect", "Cause", "Current Controls", "Recommended Action", "Owner"]
SCORE_COLS = ["Sev", "Occ", "Det", "New Sev", "New Occ", "New Det"]
FMEA_COLS = [
    "Failure Mode", "Effect", "Cause", "Current Controls", "Sev", "Occ", "Det",
    "Recommended Action", "Owner", "New Sev", "New Occ", "New Det",
]
# Columns shown in the calculated-risk table.
RESULT_COLS = ["Failure Mode", "RPN", "Risk", "Recommended Action", "Owner", "New RPN", "Reduction %"]

DEFECT_ALIASES = {
    "defect": "defect_type",
    "defect_name": "defect_type",
    "type": "defect_type",
    "failure_mode": "defect_type",
    "category": "defect_type",
    "qty": "count",
    "quantity": "count",
    "defects": "count",
    "frequency": "count",
    "number": "count",
}

# Excel fills for the risk bands (hex without '#').
EXCEL_FILL = {"High": "F8D7DA", "Medium": "FFF3CD", "Low": "D1E7DD"}


# --------------------------------------------------------------------------- defects
def clean_defects(df: pd.DataFrame) -> pd.DataFrame:
    """Validate and tidy a defect table. Returns columns defect_type, count (+ date if present)."""
    d = df.copy()
    d.columns = [str(c).strip().lower().replace(" ", "_") for c in d.columns]
    d = d.rename(columns={k: v for k, v in DEFECT_ALIASES.items() if k in d.columns and v not in d.columns})
    missing = [c for c in ("defect_type", "count") if c not in d.columns]
    if missing:
        raise ValueError(
            "missing column(s): " + ", ".join(missing) + ". Expected columns: defect_type, count (optional: date)."
        )
    d["count"] = pd.to_numeric(d["count"], errors="coerce")
    d = d.dropna(subset=["defect_type", "count"])
    d["defect_type"] = d["defect_type"].astype(str).str.strip()
    d = d[d["defect_type"].ne("") & d["count"].gt(0)]
    cols = ["defect_type", "count"]
    if "date" in d.columns:
        d["date"] = pd.to_datetime(d["date"], errors="coerce")
        cols.append("date")
    return d[cols].reset_index(drop=True)


def aggregate_defects(defects: pd.DataFrame) -> pd.DataFrame:
    """Total count per defect type."""
    if defects.empty:
        return pd.DataFrame({"defect_type": pd.Series(dtype="object"), "count": pd.Series(dtype="float64")})
    return defects.groupby("defect_type", as_index=False)["count"].sum()


def pareto_table(labels, values, cutoff: float = 80.0) -> pd.DataFrame:
    """Sort descending and add share %, cumulative % and a 'vital few' flag.

    An item is in the vital few if the cumulative share *before* adding it is still
    below the cutoff, so the item that crosses the cutoff is included.
    """
    d = pd.DataFrame({"label": list(labels), "value": list(values)})
    d = d.sort_values("value", ascending=False, kind="stable").reset_index(drop=True)
    total = d["value"].sum()
    d["share_%"] = d["value"] / total * 100 if total else 0.0
    d["cum_%"] = d["share_%"].cumsum()
    d["vital_few"] = d["cum_%"].shift(fill_value=0.0) < cutoff
    return d


# ---------------------------------------------------------------------------- FMEA
def normalize_fmea(df: pd.DataFrame) -> pd.DataFrame:
    """Make sure all FMEA columns exist, in order, with text columns as str and scores as float."""
    out = df.copy()
    for col in FMEA_COLS:
        if col not in out.columns:
            out[col] = pd.NA
    out = out[FMEA_COLS].reset_index(drop=True)
    for col in TEXT_COLS:
        s = out[col].astype("object")
        out[col] = s.where(s.notna(), "").astype(str)
    for col in SCORE_COLS:
        out[col] = pd.to_numeric(out[col], errors="coerce").astype("float64")
    return out


def merge_edit(edited: pd.DataFrame, original: pd.DataFrame) -> pd.DataFrame:
    """Restore any column the editor did not return (e.g. hidden ones) from the original rows."""
    out = edited.copy()
    for col in original.columns:
        if col not in out.columns:
            out[col] = original[col].reindex(out.index)
    return out[list(original.columns)]


def clean_fmea(df: pd.DataFrame) -> pd.DataFrame:
    """Drop rows the user added but never filled in."""
    blank = df["Failure Mode"].str.strip().eq("") & df[["Sev", "Occ", "Det"]].isna().all(axis=1)
    return df[~blank].reset_index(drop=True)


def risk_band(rpn: float, hi: float, med: float) -> str:
    if pd.isna(rpn):
        return ""
    if rpn >= hi:
        return "High"
    if rpn >= med:
        return "Medium"
    return "Low"


def compute_rpn(df: pd.DataFrame, hi: float, med: float) -> pd.DataFrame:
    """Add RPN, Risk band, New RPN and Reduction %. A blank 'New' score means 'unchanged'."""
    d = df.copy()
    d["RPN"] = d["Sev"] * d["Occ"] * d["Det"]
    new_s = d["New Sev"].fillna(d["Sev"])
    new_o = d["New Occ"].fillna(d["Occ"])
    new_d = d["New Det"].fillna(d["Det"])
    d["New RPN"] = new_s * new_o * new_d
    d["Reduction %"] = ((d["RPN"] - d["New RPN"]) / d["RPN"] * 100).where(d["RPN"] > 0)
    d["Risk"] = d["RPN"].map(lambda v: risk_band(v, hi, med))
    return d


def summary_stats(calc: pd.DataFrame, hi: float) -> dict | None:
    """Headline numbers for the Impact tab. None when there is nothing scored yet."""
    s = calc.dropna(subset=["RPN"])
    if s.empty or s["RPN"].sum() <= 0:
        return None
    before = float(s["RPN"].sum())
    after = float(s["New RPN"].sum())
    n_top = min(3, len(s))
    return {
        "before": before,
        "after": after,
        "reduction_pct": (before - after) / before * 100,
        "high_before": int((s["RPN"] >= hi).sum()),
        "high_after": int((s["New RPN"] >= hi).sum()),
        "n_top": n_top,
        "top_share_pct": float(s["RPN"].nlargest(n_top).sum()) / before * 100,
    }


def headline(stats: dict) -> str:
    return (
        f"Top {stats['n_top']} failure modes drive {stats['top_share_pct']:.0f}% of total RPN; "
        f"planned actions cut total RPN from {stats['before']:.0f} to {stats['after']:.0f} "
        f"({stats['reduction_pct']:.0f}% reduction)."
    )


def freq_vs_risk(agg: pd.DataFrame, calc: pd.DataFrame) -> pd.DataFrame:
    """Join defect frequency to FMEA risk by name (case-insensitive)."""
    s = calc.dropna(subset=["RPN"]).copy()
    s["key"] = s["Failure Mode"].str.strip().str.lower()
    a = agg.copy()
    a["key"] = a["defect_type"].astype(str).str.strip().str.lower()
    merged = s.merge(a[["key", "count"]], on="key", how="inner").drop(columns="key")
    total = agg["count"].sum()
    merged["Share of defects %"] = merged["count"] / total * 100 if total else 0.0
    return merged


# --------------------------------------------------------------------------- export
def fmea_to_csv(df: pd.DataFrame) -> bytes:
    return df[FMEA_COLS].to_csv(index=False).encode("utf-8")


def build_excel(
    defects: pd.DataFrame,
    agg: pd.DataFrame,
    calc: pd.DataFrame,
    hi: float,
    med: float,
    cutoff: float = 80.0,
) -> bytes:
    """Build the Excel report (Summary, Defects, Pareto, FMEA) and return it as bytes."""
    from openpyxl.styles import Alignment, Font, PatternFill

    stats = summary_stats(calc, hi)
    total_defects = float(agg["count"].sum()) if not agg.empty else 0.0
    summary_rows = [
        ("Total defects", int(total_defects) if total_defects else 0),
        ("Defect types", int(len(agg))),
    ]
    if stats:
        summary_rows += [
            ("Total RPN before", round(stats["before"])),
            ("Total RPN after", round(stats["after"])),
            ("RPN reduction %", round(stats["reduction_pct"], 1)),
            ("High-risk failure modes (before -> after)", f"{stats['high_before']} -> {stats['high_after']}"),
            ("Summary", headline(stats)),
        ]
    summary_rows.append(("Risk bands", f"High >= {hi:g}, Medium >= {med:g}, otherwise Low"))
    summary = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    defects_out = defects.copy()
    if "date" in defects_out.columns:
        defects_out["date"] = defects_out["date"].dt.date

    if agg.empty:
        pareto_out = pd.DataFrame(columns=["Defect", "Count", "Share %", "Cumulative %", "Vital few"])
    else:
        t = pareto_table(agg["defect_type"], agg["count"], cutoff)
        pareto_out = pd.DataFrame(
            {
                "Defect": t["label"],
                "Count": t["value"],
                "Share %": t["share_%"].round(1),
                "Cumulative %": t["cum_%"].round(1),
                "Vital few": t["vital_few"].map({True: "Yes", False: "No"}),
            }
        )

    fmea_cols = [
        "Failure Mode", "Effect", "Cause", "Current Controls", "Sev", "Occ", "Det", "RPN", "Risk",
        "Recommended Action", "Owner", "New Sev", "New Occ", "New Det", "New RPN", "Reduction %",
    ]
    fmea_out = calc.reindex(columns=fmea_cols)
    fmea_out["Reduction %"] = fmea_out["Reduction %"].round(1)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        defects_out.to_excel(xw, sheet_name="Defects", index=False)
        pareto_out.to_excel(xw, sheet_name="Pareto", index=False)
        fmea_out.to_excel(xw, sheet_name="FMEA", index=False)

        header_fill = PatternFill("solid", start_color="1F3A5F", end_color="1F3A5F")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            for column in ws.columns:
                longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                ws.column_dimensions[column[0].column_letter].width = min(max(10, longest + 2), 60)

        ws = xw.book["FMEA"]
        headers = [c.value for c in ws[1]]
        for name in ("RPN", "New RPN"):
            col_idx = headers.index(name) + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    band = risk_band(cell.value, hi, med)
                    cell.fill = PatternFill("solid", start_color=EXCEL_FILL[band], end_color=EXCEL_FILL[band])
    return buf.getvalue()
